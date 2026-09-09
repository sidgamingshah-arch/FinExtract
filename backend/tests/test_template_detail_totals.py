"""A calculated total is a real, selectable node on the Template screen — and a configuration that
cannot be validated is refused.

Defects of one family are pinned here: a surface that answers confidently about the wrong thing, and
a gate that degrades to silence instead of failing.

"RULEBOOK" HERE IS A LINE-ITEM SET. The detail serves it as the ``line_items`` block (it was
``ontology``, naming an ``ontology_versions`` row), the store is ``line_item_versions``, the door is
``POST /line-items`` and the inline edit is ``PATCH /line-items/versions/{id}/items`` keyed ``key``
(it was ``canonical_key``). Line items is the single configuration engine; there is no second
configuration surface for a screen to describe.

THE TREE WALK. ``get_template_detail`` used to treat every entry in a statement's ``sections[]`` as
a heading and look for line items only among that entry's ``children``. The shipped template
deliberately declares its calculated lines as CHILDLESS top-level sections — Gross Profit sits
between cost of sales and operating expenses because that is where the statement prints it — so each
of them came out as an inert heading row with no ``node_config`` entry, and the screen resolved the
click to whichever concept its fallback chain reached. An analyst could then edit Property, Plant and
Equipment's aliases under a heading reading "Gross profit". The walk now branches on ``role``, the
way ``services.export._emit_nodes`` does.

THE PUBLISH GATE. ``_publish_new_version`` validated an edit against the target template only ``if
tpl_row is not None``, so a configuration whose ``target_template_key`` matched no stored template
published unvalidated and became the one in force while mapping onto keys no template declares.

WHETHER A LINE IS MAPPED. Serving the calculated totals as selectable lines widened the class of
node the editor opens, including onto keys the configuration in force has no line item for: the
editor opened fully enabled and Save came back 404. ``node_config`` carries ``mapped``, and the test
below holds the flag to the configuration itself rather than to a transcribed list of keys.

A SPACER IS NOT A LINE. ``LineRole.SPACER`` is a presentational gap. The detail walk and
``export._emit_nodes`` both tested ``role == "header"`` exactly, so a spacer fell through to the
figure branch of each: a selectable concept on screen, and a row in the workbook with a label, a
value column per period and a canonical key an extracted figure could attach to. No shipped template
declares one, so it is pinned with a template published here that does.

Order is asserted against the shipped file rather than a transcribed list: a second spelling of the
template's own order is a second thing that has to keep agreeing with it.
"""
from __future__ import annotations

import json
from pathlib import Path

import pytest

API = "/api/v1"
SEEDED_TEMPLATE = "hkfrs_hk_china_v1"
# The template the ONE shipped configuration targets. ``hkfrs_hk_china_v1`` above is seeded
# TEMPLATE-ONLY now (see ``app.sample.reference``): there is a single configuration and it targets
# this key, so any test whose subject is the set IN FORCE has to ask about this pair. The tree/order
# tests keep asking about the HKFRS template, which is what ``TEMPLATE`` below is loaded from.
CONFIGURED_TEMPLATE = "output_csv_hk_v1"

_DIR = Path(__file__).resolve().parent.parent / "app" / "sample" / "templates"
TEMPLATE = json.loads((_DIR / "hkfrs_hk_china_template.json").read_text(encoding="utf-8"))


def _statement(stype: str) -> dict:
    return next(s for s in TEMPLATE["statements"] if s["type"] == stype)


def _row_id(node: dict) -> str:
    """The tree id the endpoint gives this node: headings are addressed by section, lines by key."""
    return f"sec:{node['node_id']}" if node.get("role") == "header" else node["canonical_key"]


def _statement_level_lines() -> list[str]:
    """Every line the template prints at statement level — the childless non-heading sections.

    Mostly the calculated totals (Gross profit, Net assets, Closing cash), plus the odd extracted
    line the statement prints outside any section. Derived from the file so a line added to a future
    template is covered without editing a list here: the whole CLASS of node has to stay selectable,
    not just the seventeen shipped today.
    """
    return [sec["canonical_key"]
            for stmt in TEMPLATE["statements"]
            for sec in stmt.get("sections", [])
            if sec.get("role") != "header" and not sec.get("children")]


def _template_line_items() -> int:
    """The nodes of the shipped template that CARRY A FIGURE — the count ``POST /templates`` reports.

    Spelled independently of the endpoint (which asks ``review_lines.is_statement_line``) so this
    stays a check and not a restatement. A heading and a spacer are captions; everything else is a
    line. The shipped template declares no spacer, so the number is the same either way.
    """
    from app.schemas.loader import load_template

    return len([n for n in load_template(TEMPLATE).all_nodes()
                if n.role.value not in ("header", "spacer")])


def _detail(client, locale: str = "en", template_key: str = SEEDED_TEMPLATE) -> dict:
    tpl = max((r for r in client.get(f"{API}/templates").json()
               if r["template_key"] == template_key), key=lambda r: r["version"])
    r = client.get(f"{API}/templates/{tpl['id']}/detail?locale={locale}")
    assert r.status_code == 200, r.text
    return r.json()


def _configured_detail(client) -> dict:
    """The detail for the pair that HAS a configuration in force."""
    return _detail(client, template_key=CONFIGURED_TEMPLATE)


def _statement_rows(tree: list[dict], stype: str) -> list[dict]:
    """The rows between this statement's own heading and the next one."""
    start = next(i for i, n in enumerate(tree) if n["id"] == f"stmt:{stype}")
    rest = tree[start + 1:]
    end = next((i for i, n in enumerate(rest) if str(n["id"]).startswith("stmt:")), len(rest))
    return rest[:end]


# --- a calculated total is a selectable line -------------------------------------------------

def test_gross_profit_is_a_selectable_line_not_a_heading(client):
    """The row the analyst clicks and the concept the editor opens must be the same thing."""
    detail = _detail(client)
    row = next((n for n in detail["tree"] if n["id"] == "pl_gross_profit"), None)
    assert row is not None, (
        "Gross profit reached the tree only as 'sec:pl_gross_profit' — a heading id no "
        "node_config key can match, which is what sent the click to another concept")
    assert not row.get("head"), "a calculated total carries a figure; it is a line, not a heading"
    cfg = detail["node_config"].get("pl_gross_profit")
    assert cfg is not None, "a selectable line must carry the rules the editor edits"
    assert cfg["canonical_key"] == "pl_gross_profit"
    assert cfg["label"] == "Gross profit"
    # A total printed at statement level has no section above it, so it breadcrumbs to the
    # statement rather than borrowing the heading of whichever section happens to precede it.
    assert cfg["breadcrumb"] == "Profit And Loss"
    assert cfg["aggregation"] == "Sum of children"


@pytest.mark.parametrize("key", _statement_level_lines())
def test_every_statement_level_line_is_editable(client, key):
    """Not just Gross Profit: the whole class, in all three statements."""
    detail = _detail(client)
    assert key in detail["node_config"], f"{key} is a line the screen cannot answer about"
    row = next((n for n in detail["tree"] if n["id"] == key), None)
    assert row is not None and not row.get("head")


def test_a_totals_label_localizes_like_any_other_line(client):
    """A total reaches `_loc` on the same path as a line — it must not fall back to English."""
    detail = _detail(client, locale="zh")
    zh = _statement("profit_and_loss")["sections"]
    expected = next(s for s in zh if s["canonical_key"] == "pl_gross_profit")["label_i18n"]["zh"]
    assert detail["node_config"]["pl_gross_profit"]["label"] == expected


# --- position: the user's item 4, held from the Template-screen side --------------------------

def test_the_pl_tree_is_in_the_templates_own_order(client):
    """Serving a total as a real node must not move it. The template's order IS the statement's
    order, and a calculated line printed mid-statement that renders at the end is a spread no
    analyst can read against the filing."""
    rows = _statement_rows(_detail(client)["tree"], "profit_and_loss")
    top = [n["id"] for n in rows if n["lvl"] == 1]
    assert top == [_row_id(s) for s in _statement("profit_and_loss")["sections"]]


def test_gross_profit_sits_between_cost_of_sales_and_operating_expenses(client):
    """The concrete placement, named, because "in file order" is only reassuring if you know what
    the file says: Gross profit follows the cost-of-sales block and precedes operating expenses.

    This is the reviewer's positioning requirement stated as an assertion — "gross profit should
    come [after] cost of sales on P&L". It reads the cost lines directly now that the intermediate
    total-cost-of-sales subtotal is retired, so there is nothing between the last cost line and the
    margin it produces.
    """
    ids = [n["id"] for n in _statement_rows(_detail(client)["tree"], "profit_and_loss")]
    assert (ids.index("sec:pl_s2a_cost_of_sales")
            < ids.index("pl_expenses__cost_of_goods_sold")
            < ids.index("pl_expenses__purchases_of_stock_in_trade")
            < ids.index("pl_gross_profit")
            < ids.index("sec:pl_s2_expenses")
            < ids.index("pl_expenses__taxes_and_surcharges")
            < ids.index("pl_expenses__total_operating_cost")
            < ids.index("pl_operating_profit_ebit"))


# --- genuine headings keep behaving exactly as they did --------------------------------------

def test_a_genuine_heading_is_still_a_heading_over_its_children(client):
    detail = _detail(client)
    tree = detail["tree"]
    at = next(i for i, n in enumerate(tree) if n["id"] == "sec:pl_s1_income")
    assert tree[at]["head"] is True and tree[at]["lvl"] == 1
    assert tree[at]["id"] not in detail["node_config"], (
        "a heading carries no figure and no matching rules, so it must not be selectable")

    section = next(s for s in _statement("profit_and_loss")["sections"]
                   if s["node_id"] == "pl_s1_income")
    children = tree[at + 1:at + 1 + len(section["children"])]
    assert [c["id"] for c in children] == [c["canonical_key"] for c in section["children"]]
    assert all(c["lvl"] == 2 and not c.get("head") for c in children)
    assert all(c["id"] in detail["node_config"] for c in children)


def test_the_line_item_count_is_the_templates_real_one(client):
    """One quantity, one spelling: the screen's count and the count `POST /templates` reports on
    upload are the same number. They disagreed while the seventeen statement-level lines were
    counted as headings here (170) and as line items there (187)."""
    detail = _detail(client)
    expected = _template_line_items()
    assert detail["template"]["line_items"] == expected
    assert len(detail["node_config"]) == expected
    lines = [n for n in detail["tree"]
             if not n.get("head") and not str(n["id"]).startswith("stmt:")]
    assert len(lines) == expected


# --- what the configuration in force maps ----------------------------------------------------

def _declared_keys(client, detail: dict) -> set[str]:
    """The keys the configuration IN FORCE declares, read from that configuration.

    Read back through the API rather than off the shipped file: ``mapped`` is a claim about whichever
    set ``config_select.select_for_template`` chose for THIS database, so a test transcribing the
    file would keep agreeing with itself while the screen described another set's items.

    The block is ``detail["line_items"]`` and the download is ``GET /line-items/versions/{id}``;
    both were the ``ontology`` block and ``GET /ontologies/{id}``, and the keys were under
    ``definition["mappings"][*]["canonical_key"]`` rather than ``definition["items"][*]["key"]``.
    """
    cfg = detail["line_items"]
    assert cfg, "the configured template should have a line-item set in force"
    r = client.get(f"{API}/line-items/versions/{cfg['id']}")
    assert r.status_code == 200, r.text
    # ``namespace`` filter as in ``get_template_detail``: the off-template entries are parts OF a
    # line, never template lines, so they are not keys this screen can be asked about.
    return {i.get("key") for i in (r.json()["definition"].get("items") or [])
            if i.get("namespace", "template") == "template"}


def test_every_line_says_whether_the_configuration_in_force_maps_it(client):
    """The editor opens off ``node_config``, so ``node_config`` is where the answer has to be.

    Without it the screen cannot tell an editable line from one whose every write the server
    refuses, and it offered the same fully-enabled editor for both.

    Asked of the CONFIGURED pair. The old spelling asked it of ``hkfrs_hk_china_v1``, which had an
    ontology in force; that template is seeded template-only now, so the question there is answered
    by the companion test below and the flag itself is exercised here.
    """
    detail = _configured_detail(client)
    silent = [k for k, c in detail["node_config"].items() if "mapped" not in c]
    assert not silent, (
        f"{len(silent)} lines say nothing about whether the configuration maps them (e.g. "
        f"{sorted(silent)[:3]}) — the screen cannot know which of its controls would be refused")
    declared = _declared_keys(client, detail)
    wrong = sorted(k for k, c in detail["node_config"].items() if c["mapped"] != (k in declared))
    assert not wrong, f"`mapped` disagrees with the configuration in force for {wrong}"
    # THE SHIPPED PAIR IS FULLY MAPPED, and that is the merge's own measurement rather than an
    # accident: the projection that produced this configuration placed 462 of 462 template concepts
    # with 0 fields homeless. The old assertion here was the opposite one — "a `mapped` that is True
    # everywhere is the flag not being computed at all" — and it held of the ontology in force, which
    # had no concept for ``bs_liabilities__total_liabilities``. Keeping it would pin a coverage HOLE
    # as though it were the invariant. What is asserted instead is the fact that replaced it, and the
    # flag is proved to be computed by the companion test below, where it comes back False.
    assert all(c["mapped"] for c in detail["node_config"].values()), sorted(
        k for k, c in detail["node_config"].items() if not c["mapped"])[:10]


_UNCONFIGURED_TPL_KEY = "unconfigured_probe_tpl"
# HEADERS ONLY, and that is now the whole point. A template carrying any line that holds a figure
# provisions its own configuration on upload (`routes/templates._provision_line_items`), so the
# no-configuration state can no longer be reached by publishing an ordinary template — the
# previous version of this probe had one `role: "line"` and stopped being unconfigured the moment
# that rule landed. A section header names a section rather than a figure and is deliberately not
# provisioned, so a header-only template is the one template that legitimately has nothing in
# force. Which is the honest shape for this test anyway: it asks what the screen says when there
# is no configuration, and this is now the only way there isn't one.
_UNCONFIGURED_TEMPLATE = {
    "template_key": _UNCONFIGURED_TPL_KEY,
    "name": "Template with no configuration",
    "statements": [{
        "type": "balance_sheet",
        "sections": [{"node_id": "bs_ca", "canonical_key": "bs_ca", "label": "Current Assets",
                      "role": "header"}],
    }],
}


def test_a_template_with_no_configuration_says_so_on_every_line(client):
    """The other side of the flag, and the proof it is computed rather than defaulted to True.

    A template with no line-item set targeting it has nothing in force, so the screen must say so —
    no ``line_items`` block, and no line reporting itself mapped. A blank the user can fix in
    configuration is the right answer; a fully-enabled editor over writes the server would refuse
    is not.

    POSTS ITS OWN TEMPLATE rather than reading a seeded key. That is not tidiness: asserting
    "nothing is published against this template" over a SHARED seeded key makes the test depend on
    the whole suite never publishing a set that targets it, and it passed alone while failing in
    the full run for exactly that reason.

    This replaces ``test_the_controls_the_screen_withholds_are_the_ones_the_server_refuses`` as the
    place the False case is pinned; see its retirement note below.
    """
    r = client.post(f"{API}/templates", json={"definition": _UNCONFIGURED_TEMPLATE})
    assert r.status_code == 201, r.text
    # Nothing was provisioned, and the upload says so by omitting the block rather than by
    # reporting an empty one.
    assert "line_item_version" not in r.json(), (
        "a header-only template implies no lines, so it must provision no configuration")

    detail = _detail(client, template_key=_UNCONFIGURED_TPL_KEY)
    assert detail["line_items"] is None, (
        "nothing is published against this template, so the screen must not name a configuration")
    assert not any(c["mapped"] for c in (detail["node_config"] or {}).values()), (
        "no set targets this template, so no line of it is mapped")


def test_an_ordinary_template_arrives_configured(client):
    """The counterpart, and the rule that changed the probe above.

    Any template with a figure-bearing line provisions its configuration on upload, so `mapped` is
    True on every one of its lines from the moment it exists. Before this, a fresh template served
    every line unmapped and an author had to write the whole configuration by hand first.
    """
    definition = {
        "template_key": "configured_on_upload_tpl",
        "name": "Configured on upload",
        "statements": [{
            "type": "balance_sheet",
            "sections": [{
                "node_id": "bs_ca", "canonical_key": "bs_ca", "label": "Current Assets",
                "role": "header",
                "children": [{"node_id": "bs_ca__cash", "canonical_key": "bs_ca__cash",
                              "label": "Cash", "role": "line"}],
            }],
        }],
    }
    r = client.post(f"{API}/templates", json={"definition": definition})
    assert r.status_code == 201, r.text
    assert r.json().get("line_item_version"), "the template provisioned no configuration"

    detail = _detail(client, template_key="configured_on_upload_tpl")
    assert detail["line_items"] is not None, "the screen names no configuration for it"
    assert detail["node_config"]["bs_ca__cash"]["mapped"] is True


# RETIRED: test_the_calculated_total_this_round_exposed_is_the_unmapped_one.
#
# Its whole subject was that the rulebook in force for ``hkfrs_hk_china_v1`` had no concept for
# ``bs_liabilities__total_liabilities`` while it did have one for the other sixteen statement-level
# lines. Both halves of that premise are gone: nothing is published against that template any more,
# and the one shipped configuration covers its target template completely (462/462, 0 homeless), so
# there is no "the unmapped one" to name. The CLASS is still covered — every statement-level line is
# selectable (``test_every_statement_level_line_is_editable``) and every line states whether it is
# mapped (the two tests above, one for each answer).

_UNRESOLVABLE_TPL_KEY = "unresolvable_probe_tpl"
_UNRESOLVABLE_CFG_KEY = "unresolvable_probe_cfg"
_UNRESOLVABLE_TEMPLATE = {
    "template_key": _UNRESOLVABLE_TPL_KEY,
    "name": "Unresolvable configuration probe",
    "statements": [{
        "type": "balance_sheet",
        "sections": [{"node_id": "cash", "canonical_key": "probe_cash", "label": "Cash",
                      "role": "line"}],
    }],
}
_UNRESOLVABLE_SET = {
    "line_items_key": _UNRESOLVABLE_CFG_KEY,
    "target_template_key": _UNRESOLVABLE_TPL_KEY,
    "schema_version": 1,
    # A section layer, so the fold runs at all…
    "section_defaults": {"bs_s1": {"statement": "balance_sheet", "section_scope": ["bs_s1"]}},
    # …and an `inherits` naming an entry that is not in it. The definition VALIDATES, and resolving
    # the section layer raises on it — so `get_template_detail`'s
    # `load_line_item_set(resolve=True)` fails and its `except` serves the screen no rules at all.
    # The ITEM is declared regardless, and the declaration is what the editing endpoint looks for.
    "items": [{"key": "probe_cash", "label": "Cash", "aliases": ["Cash"],
               "inherits": "no_such_section"}],
}


@pytest.fixture
def unresolvable_config(client):
    """A stored configuration the detail cannot LOAD, targeting a template of one line.

    Inserted straight into the database because ``POST /line-items`` refuses an unresolvable
    ``inherits`` (422) — a row like this predates that gate, which is the state the ``except`` in
    ``get_template_detail`` exists for. It was an ``ontology_versions`` row seeded the same way;
    line items is the single configuration engine, so the store is ``line_item_versions``.
    """
    from app.db.models import LineItemVersion, TemplateVersion

    from app.db.base import SessionLocal

    r = client.post(f"{API}/templates", json={"definition": _UNRESOLVABLE_TEMPLATE})
    assert r.status_code == 201, r.text
    with SessionLocal() as session:
        session.add(LineItemVersion(line_items_key=_UNRESOLVABLE_CFG_KEY,
                                    target_template_key=_UNRESOLVABLE_TPL_KEY,
                                    version=1, definition=_UNRESOLVABLE_SET))
        session.commit()
    try:
        yield r.json()
    finally:
        _drop(LineItemVersion, LineItemVersion.line_items_key, _UNRESOLVABLE_CFG_KEY)
        _drop(TemplateVersion, TemplateVersion.template_key, _UNRESOLVABLE_TPL_KEY)


def test_a_configuration_that_cannot_be_loaded_still_says_which_items_it_declares(
        client, unresolvable_config):
    """``mapped`` is read off the STORED definition, because that is what the writes are checked on.

    Taken from the loaded configuration, one unresolvable ``inherits`` — and 475 of 475 shipped items
    inherit their gate — turned the flag inside out: the screen locked every line and said the
    configuration declared no item for any of them, while the edit endpoint went on accepting the key
    it had just disabled.
    """
    r = client.get(f"{API}/templates/{unresolvable_config['id']}/detail")
    assert r.status_code == 200, r.text
    detail = r.json()
    assert detail["node_config"]["probe_cash"]["mapped"] is True, (
        "the configuration declares this item; only its RULES could not be loaded")

    # THE SECOND HALF CHANGED, AND ON PURPOSE. On the retired ontology route this edit was ACCEPTED
    # (200) and the assertion was "the write `mapped` is a claim about" — mapped True meant the write
    # went through. ``_publish_new_version`` loads with ``resolve=True`` now, so it refuses to
    # publish another version that cannot be resolved, and the refusal NAMES the dangling `inherits`
    # rather than pretending the key is unknown. That is the honest answer for this row: the item IS
    # declared, and the thing wrong with the configuration is the section it points at.
    #
    # The invariant this test exists for is untouched: ``mapped`` is read off the STORED definition,
    # so one dangling `inherits` cannot make the screen say the configuration declares no item for
    # any of its lines. A 404 here would be that inversion reappearing on the write side.
    r = client.patch(f"{API}/line-items/versions/{detail['line_items']['id']}/items",
                     json={"key": "probe_cash", "aliases": ["Cash at bank"]})
    assert r.status_code == 422, r.text
    body = json.dumps(r.json())
    assert "no_such_section" in body and "probe_cash" in body, (
        f"the refusal has to name the section it could not find, or the author cannot fix it: {body}")


# --- an unvalidatable configuration is refused ------------------------------------------------

_PROBE_KEY = "orphan_probe_tpl"
_PROBE_CFG_KEY = "orphan_probe_cfg"
_PROBE_TEMPLATE = {
    "template_key": _PROBE_KEY,
    "name": "Orphan probe",
    "statements": [{
        "type": "balance_sheet",
        "sections": [
            {"node_id": "cash", "canonical_key": "cash", "label": "Cash", "role": "line"},
            # A SECOND LINE THE CONFIGURATION BELOW DOES NOT DECLARE, so this pair carries an
            # UNMAPPED key for `test_the_controls_the_screen_withholds_are_the_ones_the_server_
            # refuses` to drive. That test used to read one off the shipped pair; the one shipped
            # configuration covers its target template completely (462/462, 0 homeless), so the
            # unmapped case has to be constructed rather than found.
            {"node_id": "inv", "canonical_key": "inv", "label": "Inventories", "role": "line"},
        ],
    }],
}
_PROBE_SET = {
    "schema_version": 1,
    "line_items_key": _PROBE_CFG_KEY,
    "target_template_key": _PROBE_KEY,
    "items": [{"key": "cash", "label": "Cash", "aliases": ["Cash"]}],
}


def _drop(model, key_column, key: str) -> None:
    from sqlalchemy import delete

    from app.db.base import SessionLocal

    with SessionLocal() as session:
        session.execute(delete(model).where(key_column == key))
        session.commit()


@pytest.fixture
def probe_pair(client):
    """A throwaway template and a configuration published against it, removed again afterwards.

    The ``client`` fixture is session-scoped and so is its database: a probe left behind would sit
    in every later test's `/templates` and `/line-items/versions` listing.
    """
    from app.db.models import LineItemVersion, TemplateVersion

    r = client.post(f"{API}/templates", json={"definition": _PROBE_TEMPLATE})
    assert r.status_code == 201, r.text
    tpl = r.json()
    r = client.post(f"{API}/line-items", json={"definition": _PROBE_SET})
    assert r.status_code == 201, r.text
    try:
        yield tpl, r.json()
    finally:
        _drop(LineItemVersion, LineItemVersion.line_items_key, _PROBE_CFG_KEY)
        _drop(TemplateVersion, TemplateVersion.template_key, _PROBE_KEY)


@pytest.fixture
def orphaned_config(probe_pair):
    """The same configuration, with its target template GONE — a renamed key, or a deleted template.

    Orphaned by dropping the template row rather than by publishing against a key that never
    existed, because the create path has always refused that: the hole was on the edit path.
    """
    from app.db.models import TemplateVersion

    _drop(TemplateVersion, TemplateVersion.template_key, _PROBE_KEY)
    return probe_pair[1]["id"]


def test_the_controls_the_screen_withholds_are_the_ones_the_server_refuses(client, probe_pair):
    """``mapped`` is not an opinion about tidiness — it is the answer to "would this be refused?".

    Both writes the screen offers for a declared key are asked here on an UNDECLARED one, so the flag
    and the server's answer cannot drift apart: the item editor's Save (404, not in this version) and
    the confusable-with picker (422, unknown line item). Neither changes stored state when refused.

    WAS THREE WRITES, ON THE SHIPPED PAIR. The third was the NETTING picker, and that editor is
    deleted with ``routes/ontologies.py``: the shipped configuration declares 0 netting rules, so it
    was an editor for an empty list (see ``routes/line_items.py``, "WHAT IS GONE"). Do not reinstate
    the third leg without code and data that read netting rules. The pair is a probe rather than the
    shipped one because the shipped configuration now covers its target template completely, so
    there is no unmapped key on it to ask about — see ``_PROBE_TEMPLATE``'s second line.
    """
    tpl, cfg = probe_pair
    detail = client.get(f"{API}/templates/{tpl['id']}/detail").json()
    cfg_id = detail["line_items"]["id"]
    unmapped = sorted(k for k, c in detail["node_config"].items() if not c["mapped"])
    assert unmapped == ["inv"], f"the probe pair must carry exactly one unmapped line: {unmapped}"
    editable = next(k for k, c in detail["node_config"].items() if c["mapped"])
    for key in unmapped:
        r = client.patch(f"{API}/line-items/versions/{cfg_id}/items",
                         json={"key": key, "aliases": ["Anything at all"]})
        assert r.status_code == 404, f"an alias edit on {key} was accepted: {r.text}"
        r = client.patch(f"{API}/line-items/versions/{cfg_id}/items",
                         json={"key": editable, "confusable_with": [key]})
        assert r.status_code == 422, f"{key} was accepted as a confusable_with target: {r.text}"


def test_an_edit_to_a_configuration_with_no_target_template_is_refused(client, orphaned_config):
    """`if tpl_row is not None` skipped validation entirely when the target template was missing,
    so this edit published — and a configuration nothing had checked became the one in force."""
    r = client.patch(f"{API}/line-items/versions/{orphaned_config}/items",
                     json={"key": "cash", "aliases": ["Cash at bank"]})
    assert r.status_code == 422, r.text
    assert _PROBE_KEY in json.dumps(r.json()), (
        "the refusal has to name the template it could not find, or the author cannot fix it")


# RETIRED: test_a_netting_edit_takes_the_same_gate.
#
# It asked the SECOND inline-edit endpoint, ``PATCH /ontologies/{id}/netting-rules``, to take the
# same target-template gate as the alias edit — "both inline-edit endpoints publish through
# `_publish_new_version`, so neither can be the one path where validation is skipped".
#
# There is only ONE inline-edit endpoint now. The netting-rules editor is deleted with
# ``routes/ontologies.py``, because it governed nothing: the shipped configuration declares 0 netting
# rules on both shipped files, so it was an editor for an empty list (``routes/line_items.py``, "WHAT
# IS GONE"). The invariant it shared — an edit whose target template is missing is refused rather
# than published unvalidated — is asserted above on the endpoint that survives, and
# ``_publish_new_version`` has no second caller left to skip it.
#
# Re-add the endpoint and this test together, and only alongside code and data that read netting
# rules.


def test_a_new_configuration_naming_no_stored_template_is_refused(client):
    # One recognising item, because the door checks "recognises anything" BEFORE it checks the target
    # template (``routes/line_items.py``); an empty `items` list would be refused for the other
    # reason and this test would stop being about the template at all.
    body = {"definition": {"schema_version": 1,
                           "line_items_key": "no_such_target_cfg",
                           "target_template_key": "template_that_was_never_published",
                           "items": [{"key": "cash", "label": "Cash", "aliases": ["Cash"]}]}}
    r = client.post(f"{API}/line-items", json=body)
    assert r.status_code == 422, r.text
    assert "template_that_was_never_published" in json.dumps(r.json())


# --- and the record says WHICH template version it was checked against -----------------------

def test_a_published_configuration_states_the_template_version_it_was_checked_against(probe_pair):
    """The check runs against whichever template version is newest at publish time, which is not
    necessarily the version a run pins. Saying so on the response is what lets a reader tell,
    rather than assume, what a configuration was held to."""
    tpl, cfg = probe_pair
    assert cfg.get("validated_against_template") == {
        "id": tpl["id"], "template_key": tpl["template_key"], "version": tpl["version"]}


def test_an_inline_edit_states_it_the_same_way(client, probe_pair):
    """Create and edit report it with one spelling, so neither path is the one you have to guess
    about."""
    tpl, cfg = probe_pair
    r = client.patch(f"{API}/line-items/versions/{cfg['id']}/items",
                     json={"key": "cash", "aliases": ["Cash at bank"]})
    assert r.status_code == 200, r.text
    against = r.json().get("validated_against_template")
    assert against == {"id": tpl["id"], "template_key": _PROBE_KEY, "version": tpl["version"]}
    assert against == cfg["validated_against_template"]


def test_an_edit_to_the_shipped_configuration_names_the_shipped_template(client):
    """Not only for a probe: the configuration actually in force reports it too.

    Asked of the CONFIGURED pair — ``hkfrs_hk_china_v1`` is template-only now, so it has nothing in
    force to edit. The key edited is one the shipped set really declares.
    """
    cfg = _configured_detail(client)["line_items"]
    assert cfg, "the configured template should have a line-item set in force"
    r = client.patch(f"{API}/line-items/versions/{cfg['id']}/items",
                     json={"key": "is_pl__gross_profit", "aliases": ["Gross profit"]})
    assert r.status_code == 200, r.text
    against = r.json().get("validated_against_template")
    assert against, "an edit that says nothing about what validated it cannot be audited"
    assert against["template_key"] == CONFIGURED_TEMPLATE
    assert against["version"] == max(r["version"] for r in client.get(f"{API}/templates").json()
                                     if r["template_key"] == CONFIGURED_TEMPLATE)
    assert against["id"]


# --- a spacer is a presentational gap, not a line ---------------------------------------------

_SPACER_KEY = "spacer_probe_tpl"
_GAP_LABEL = "— gap —"
_SPACER_TEMPLATE = {
    "template_key": _SPACER_KEY,
    "name": "Spacer probe",
    "statements": [{
        "type": "balance_sheet",
        "sections": [{
            "node_id": "ca", "canonical_key": "ca_head", "label": "Current assets",
            "role": "header",
            "children": [
                {"node_id": "cash", "canonical_key": "cash", "label": "Cash", "role": "line"},
                # Carrying a canonical_key because the schema requires one, which is exactly what
                # made the old treatment costly: the gap was a keyed row an extracted figure could
                # land on.
                {"node_id": "gap", "canonical_key": "spacer_gap", "label": _GAP_LABEL,
                 "role": "spacer"},
                {"node_id": "inv", "canonical_key": "inv", "label": "Inventories", "role": "line"},
            ],
        }],
    }],
}


@pytest.fixture
def spacer_template(client):
    """A published template that declares a spacer — no shipped one does — removed afterwards.

    Removed for the reason ``probe_pair`` is: the ``client`` fixture's database is session-scoped, so
    a probe left behind sits in every later test's `/templates` listing.
    """
    from app.db.models import TemplateVersion

    r = client.post(f"{API}/templates", json={"definition": _SPACER_TEMPLATE})
    assert r.status_code == 201, r.text
    try:
        yield r.json()
    finally:
        _drop(TemplateVersion, TemplateVersion.template_key, _SPACER_KEY)


def test_a_spacer_is_no_line_item_on_publish_or_on_the_screen(client, spacer_template):
    """A gap is not a concept: it has no aliases, no sign convention and no figure to extract.

    Counted as a line item it also breaks the one quantity `_publish` and this screen both report.
    """
    assert spacer_template["line_items"] == 2, (
        "the template declares two lines and one gap; a spacer carries no figure")
    r = client.get(f"{API}/templates/{spacer_template['id']}/detail")
    assert r.status_code == 200, r.text
    detail = r.json()
    assert detail["template"]["line_items"] == spacer_template["line_items"]
    assert sorted(detail["node_config"]) == ["cash", "inv"], (
        "a presentational gap must not be offered as a concept to alias and sign")
    ids = [n["id"] for n in detail["tree"]]
    assert ids == ["stmt:balance_sheet", "sec:ca", "cash", "sec:gap", "inv"], (
        "the gap keeps its place in the statement, addressed like a caption: `sec:gap` matches no "
        "node_config key, so clicking it cannot resolve to another concept's rules")
    gap = next(n for n in detail["tree"] if n["id"] == "sec:gap")
    assert gap["head"] is True and gap["lvl"] == 2, "a gap is not selectable"


def test_the_export_writes_a_spacer_as_a_blank_row_not_a_figure_row():
    """The workbook's equivalent of a presentational gap is an empty row, not a line item.

    Emitted through the figure branch, the gap arrived as a row with a label, a note cell and a value
    column per period — and, because it carries a canonical_key, a row `by_key` could attach an
    extracted figure to.
    """
    pytest.importorskip("openpyxl")
    import io

    import openpyxl

    from app.services.export import build_statement_workbook

    def _row(key: str, label: str, value: str) -> dict:
        return {"canonical_key": key, "source_label": label,
                "values": [{"basis": "consolidated", "period_label": "current", "value": value}]}

    rows = [_row("cash", "Cash", "10"), _row("spacer_gap", _GAP_LABEL, "4242"),
            _row("inv", "Inventories", "20")]
    wb = openpyxl.load_workbook(io.BytesIO(
        build_statement_workbook(rows, _SPACER_TEMPLATE, filename="f.pdf")))
    ws = wb["Balance Sheet"]
    col_a = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
    assert _GAP_LABEL not in col_a, f"the gap was written out as a line: {col_a}"
    at = col_a.index("Cash")
    assert col_a[at + 1] is None, f"the gap's row is the gap, so it stays empty: {col_a}"
    assert col_a[at + 2] == "Inventories", f"the lines around it keep their order: {col_a}"
    gap_row = at + 2                                   # col_a[i] is row i + 1
    assert all(ws.cell(gap_row, c).value in (None, "")
               for c in range(1, ws.max_column + 1)), "the gap row carries no cell of its own"
    written = {ws.cell(r, c).value
               for r in range(1, ws.max_row + 1) for c in range(1, ws.max_column + 1)}
    assert 4242 not in written and "4242" not in written, (
        "a figure keyed to the gap reached the sheet — a presentational row with a number in it")
