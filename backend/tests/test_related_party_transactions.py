"""Related Party Transactions is a FIGURE: the transactions the related-party note prints for the
year, added up from one part per category, with every amount and the note text carried beside it.

The user's decision, in their words: "give out a full prose as well as numbers which add up all
different transactions". Before, the line asked the model for "the complete relevant disclosure",
the model could only cite printed rows, and the export wrote Yes/No. Measured on 301577 (美信科技),
its request carried the parent-company table and the related-party GUARANTEES (100,000,000 /
200,000,000, guarantees received — not transactions) with key management pay fused in.
"""
from __future__ import annotations

import json
import pathlib
from decimal import Decimal
from types import SimpleNamespace

import pytest

from app.api.routes.extractions import NOT_DETERMINED, _build_supplemental_rows, _serialize_rows
from app.config import get_settings
from app.core.models import DocumentModel
from app.core.models.enums import Basis
from app.core.models.line_item import (
    ExtractedValue, LineItem, NoteItem, NotesTable, Provenance, UnitContext,
)
from app.core.stage import PipelineContext
from app.schemas.line_items import load_line_item_set
from app.services import line_disclosure, line_item_notes
from app.services.export import build_rows_json, disclosure_text
from app.stages.note_sourced import NoteSourcedStage

_TEMPLATES = pathlib.Path(__file__).resolve().parents[1] / "app" / "sample" / "templates"
SEED = _TEMPLATES / "output_csv_hk_line_items.json"
RPT = "notes__related_party_transactions"
PARTS = ("sub__rpt_sales", "sub__rpt_purchases", "sub__rpt_leases", "sub__rpt_interest",
         "sub__rpt_key_management", "sub__rpt_other")


@pytest.fixture(scope="module")
def shipped():
    return load_line_item_set(json.loads(SEED.read_text(encoding="utf-8")), resolve=True)


def _items(shipped):
    return {i.key: i for i in shipped.items}


# --- the configuration --------------------------------------------------------------------------

def test_the_line_is_the_sum_of_one_part_per_transaction_category(shipped):
    items = _items(shipped)
    line = items[RPT]
    assert str(line.type) == "derived"
    assert len(line.cascade) == 1
    rung = line.cascade[0]
    assert rung.terms_op == "sum"
    assert [t.ref for t in rung.terms] == list(PARTS)
    # `any_of`: a filing printing only key management pay resolves on it; one printing none of
    # the categories leaves the line empty rather than publishing zero.
    assert {t.role for t in rung.terms} == {"any_of"}
    assert {t.sign for t in rung.terms} == {1}


@pytest.mark.parametrize("key", PARTS)
def test_every_part_is_an_internal_note_part_wired_to_the_line(shipped, key):
    part = _items(shipped)[key]
    assert part.parent == RPT
    assert part.namespace == "internal" and part.in_output is False
    assert str(getattr(part.statement, "value", part.statement)) == "notes"
    assert part.route == "note_tables"
    assert part.definition.strip()
    # No face alias, so no caption can bind a part from a statement (`test_line_item_gate`).
    assert not part.aliases
    terms = set(part.note_source.note_terms)
    # Every part is pointed at the related-party CHAPTER, in both scripts and in English.
    assert {"关联方及关联交易", "related party transactions"} <= terms


def test_balances_and_guarantees_are_not_transactions(shipped):
    line = _items(shipped)[RPT]
    text = line.definition.lower()
    for word in ("balances", "guarantees", "not added"):
        assert word in text, word


def test_key_management_pay_belongs_to_its_part_not_to_the_line(shipped):
    items = _items(shipped)
    zh = set((items[RPT].aliases_i18n or {}).get("zh") or ())
    assert not zh & {"关键管理人员报酬", "关键管理人员薪酬"}, \
        "a KMP row bound to the line would be served as the whole figure"
    assert "关键管理人员报酬" in items["sub__rpt_key_management"].note_source.note_terms


# --- what the model is sent ---------------------------------------------------------------------

def _rp_notes() -> list[NotesTable]:
    """301577's shape: an other-payables note and the 十三 related-party chapter."""
    def table(number, title, chapter, caption, amount):
        row = NoteItem(raw_label=caption)
        row.values["k"] = ExtractedValue(basis=Basis.CONSOLIDATED, period_label="current",
                                         value=Decimal(amount), value_raw=Decimal(amount),
                                         provenance=Provenance(page_index=198))
        return NotesTable(note_number=number, title=title, chapter_title=chapter, items=[row],
                          source_pages=[198], source_text=f"{title} {caption} {amount}")
    return [
        table("七、21", "、其他应付款", "合并财务报表项目注释", "其他应付款", "5000000.00"),
        table("十三、1", "、本企业的母公司情况", "关联方及关联交易", "深圳全珍投资", "30000000"),
        table("十三、4", "关联担保情况", "关联方及关联交易", "张定珍、胡联全", "100000000.00"),
        table("十三、4", "关键管理人员报酬", "关联方及关联交易", "关键管理人员薪酬", "4068320.71"),
    ]


@pytest.mark.parametrize("key", PARTS)
def test_each_part_is_sent_the_related_party_chapter_and_not_other_payables(shipped, key):
    part = _items(shipped)[key]
    hits = line_item_notes.notes_for_line_item(part, line_item_notes.header_pool(_rp_notes()))
    got = [h.note for h in hits]
    assert "十三、4" in got, got
    assert "七、21" not in got, got


# --- the figure, its items and its prose --------------------------------------------------------

def _llm_part(key: str, label: str, amount: str, note: str, page: int) -> LineItem:
    """A part as the model route leaves it: a figure and the trail of the row it cited."""
    row = LineItem(source_label=key, canonical_key=key)
    row.values["c"] = ExtractedValue(basis=Basis.CONSOLIDATED, period_label="current",
                                     value=Decimal(amount), value_raw=Decimal(amount))
    row.derivation = {"consolidated:current": {
        "method": "line_item_llm:1_cited_row", "result": amount, "version": 2, "op": "first",
        "inputs": [{"label": label, "value": amount, "note": note, "counted": True,
                    "status": "used", "used": amount,
                    "provenance": {"page_index": page}}]}}
    return row


def _run(shipped, rows):
    doc = DocumentModel(filename="f.pdf")
    doc.notes = _rp_notes()
    doc.line_items = list(rows)
    ctx = PipelineContext(settings=get_settings())
    ctx.line_items = shipped
    return NoteSourcedStage().run(doc, ctx)


def _row(doc, key):
    return next((li for li in doc.line_items if li.canonical_key == key), None)


def test_the_line_adds_its_parts_and_carries_every_amount_and_the_note_text(shipped):
    sales = _llm_part("sub__rpt_sales", "销售商品", "2046232.67", "十三、4", 198)
    doc = _run(shipped, [sales])
    line = _row(doc, RPT)
    assert line is not None
    figure = next(ev.value for ev in line.values.values() if ev.period_label == "current")
    # 2,046,232.67 of sales, read by the model, and 4,068,320.71 of key management pay, read by
    # the note route off the 关键管理人员报酬 table — the guarantee's 100,000,000 is in neither.
    assert figure == Decimal("2046232.67") + Decimal("4068320.71")
    items = line.disclosure["items"]
    assert {(i["category"], i["amount"]) for i in items if i["period"] == "current"} == {
        ("Sales of goods and services to related parties", "2046232.67"),
        ("Key management personnel compensation", "4068320.71")}
    sale = next(i for i in items if i["part"] == "sub__rpt_sales")
    assert sale["caption"] == "销售商品" and sale["note"] == "十三、4"
    assert sale["provenance"]["page_index"] == 198
    # The chapter's text, quoted — the guarantees and the parent company included, the other
    # payables note not.
    quoted = {(n["note"], n["title"]) for n in line.disclosure["notes"]}
    assert ("十三、4", "关联担保情况") in quoted and ("十三、1", "、本企业的母公司情况") in quoted
    assert not any(n == "七、21" for n, _ in quoted)
    guarantee = next(n for n in line.disclosure["notes"] if n["title"] == "关联担保情况")
    assert guarantee["text"] == "关联担保情况 张定珍、胡联全 100000000.00"


def test_a_line_with_no_part_resolved_carries_nothing(shipped):
    doc = DocumentModel(filename="f.pdf")
    doc.notes = [_rp_notes()[0]]
    ctx = PipelineContext(settings=get_settings())
    ctx.line_items = shipped
    NoteSourcedStage().run(doc, ctx)
    assert _row(doc, RPT) is None


def test_only_a_printed_notes_line_built_from_note_parts_carries_a_disclosure(shipped):
    lines = line_disclosure.disclosing_lines(shipped.items)
    assert RPT in lines
    assert [p.key for p in lines[RPT]] == list(PARTS)
    # Other receivables' net is built from note parts too, but it is a step inside another line.
    assert "sub__cp_other_receivables_net" not in lines
    assert all(k.startswith("notes__") for k in lines)


def test_the_served_row_and_the_export_carry_the_items_and_the_quote(shipped):
    doc = _run(shipped, [_llm_part("sub__rpt_sales", "销售商品", "2046232.67", "十三、4", 198)])
    rows = _serialize_rows(doc)
    served = next(r for r in rows if r["canonical_key"] == RPT)
    assert served["disclosure"]["items"]
    text = disclosure_text(served)
    assert "Sales of goods and services to related parties — 销售商品 (note 十三、4, p.199): 2046232.67" in text
    assert "As the filing prints it:" in text and "张定珍、胡联全 100000000.00" in text
    exported = json.loads(build_rows_json(rows, filename="f.pdf"))
    line = next(r for r in exported["line_items"] if r["canonical_key"] == RPT)
    assert line["disclosure"] == served["disclosure"]


def test_the_line_items_sheet_writes_the_figure_and_the_disclosure_in_its_comment(shipped):
    from io import BytesIO

    from openpyxl import load_workbook

    from app.services.export import build_statement_workbook
    from app.services.export_line_items import SHEET

    doc = _run(shipped, [_llm_part("sub__rpt_sales", "销售商品", "2046232.67", "十三、4", 198)])
    rows = _serialize_rows(doc)
    tpl = json.loads((_TEMPLATES / "output_csv_hk_v1_template.json").read_text(encoding="utf-8"))
    wb = load_workbook(BytesIO(build_statement_workbook(rows, tpl, line_item_set=shipped)))
    ws = wb[SHEET]
    hit = next(row for row in ws.iter_rows() if row[1].value == RPT)
    # The figure, never "Yes": a live sum over the two parts' rows, which hold what they read.
    formula = hit[4].value
    assert isinstance(formula, str) and formula.startswith("="), formula
    by_cell = {f"E{row[0].row}": row[4].value for row in ws.iter_rows()}
    added = sorted(Decimal(str(by_cell[ref])) for ref in formula[1:].split("+"))
    assert added == [Decimal("2046232.67"), Decimal("4068320.71")]
    assert "Added up from 2 amounts" in hit[0].comment.text
    assert "关联担保情况" in hit[0].comment.text


# --- the supplemental row ----------------------------------------------------------------------

def test_the_presence_scan_no_longer_answers_yes():
    tpl = json.loads((_TEMPLATES / "output_csv_hk_v1_template.json").read_text(encoding="utf-8"))
    doc = SimpleNamespace(unit_context=UnitContext(currency="CNY", scale_factor=Decimal(1)),
                          notes=[], fmt=None, pages=[])
    out = _build_supplemental_rows(
        template_def=tpl, base_rows=[], entity_name="X", doc_model=doc, ctx=None, options={},
        disclosures=[{"key": "related_party", "present": True, "page": 41}])
    row = next(r for r in out if r["canonical_key"] == RPT)
    assert row["values"] == []
    assert NOT_DETERMINED in row["flags"]


# --- the reader: two related-party tables under one heading --------------------------------------

def test_a_second_counterparty_table_is_read_under_its_own_header():
    """002273 p185: 采购商品/接受劳务情况表 (six columns, 获批的交易额度 among them) and then
    出售商品/提供劳务情况表 (four), under one heading "（1）购销商品、提供和接受劳务的关联交易". Read as one
    section the sales took the purchases' grid and every sale was filed under 获批的交易额度."""
    from app.core.models.geometry import BBox
    from app.services.notes_extract import extract_note_tables
    from app.services.row_reconstruct import Word

    fix = json.loads((pathlib.Path(__file__).parent / "fixtures/related_party_tables.json")
                     .read_text(encoding="utf-8"))["002273_rp_sales"]
    words = [Word(text=t, bbox=BBox(x0=a, y0=b, x1=c, y1=d)) for t, a, b, c, d in fix["words"]]
    tables = extract_note_tables(words, page_index=fix["page"], document_id="t",
                                 source_kind="native", carry_note=(fix["note"], fix["title"]))
    by_title = {t.title: t for t in tables}
    assert set(by_title) == {"购销商品、提供和接受劳务的关联交易", "出售商品/提供劳务情况表"}

    def figures(title, caption):
        row = next(i for i in by_title[title].items if i.raw_label == caption)
        return {v.period_label: str(v.value) for v in row.values.values()}

    assert figures("出售商品/提供劳务情况表", "重庆晶朗光电有限公司销售商品") == {"prior": "56476.40"}
    assert figures("出售商品/提供劳务情况表", "东莞埃科思科技有限公司销售设备") == {"current": "966.91"}
    assert figures("购销商品、提供和接受劳务的关联交易", "光驰公司采购设备") == {
        "current": "34228256.61", "prior": "118964045.72"}


def test_a_row_printed_in_the_inner_column_takes_the_nature_printed_above_it():
    """0285 note 35: "Nature of transactions | Notes | Related parties". The nature is printed once
    per pair and the second row says only "Fellow subsidiaries" — ten such rows under one group, so
    no citation could tell 17,989,310 of sales from 13,170 of plant purchases."""
    from app.core.models.geometry import BBox
    from app.services.notes_extract import extract_note_tables
    from app.services.row_reconstruct import Word

    fix = json.loads((pathlib.Path(__file__).parent / "fixtures/related_party_tables.json")
                     .read_text(encoding="utf-8"))["0285_rp_nature"]
    words = [Word(text=t, bbox=BBox(x0=a, y0=b, x1=c, y1=d)) for t, a, b, c, d in fix["words"]]
    tables = extract_note_tables(words, page_index=fix["page"], document_id="t",
                                 source_kind="native", carry_note=(fix["note"], fix["title"]))
    rows = [i for t in tables for i in t.items if i.raw_label == "Fellow subsidiaries"]
    groups = {i.group_hint: {v.period_label: str(v.value) for v in i.values.values()} for i in rows}
    assert len(groups) == len(rows), [i.group_hint for i in rows]
    sales = next(g for g in groups if g.startswith("Sales of inventories"))
    assert groups[sales] == {"current": "17989310", "prior": "11536290"}
    plant = next(g for g in groups if g.startswith("Purchases of plant and machinery"))
    assert groups[plant] == {"current": "13170", "prior": "24685"}


def test_a_counterparty_wrapped_round_its_row_names_the_row():
    """300391 p245: each counterparty's name is printed over three lines, the figures' line carrying
    only 销售商品 — four rows with one caption and four different figures, which no citation could
    tell apart. Each takes the name printed beside it."""
    from app.core.models.geometry import BBox
    from app.services.notes_extract import extract_note_tables
    from app.services.row_reconstruct import Word

    fix = json.loads((pathlib.Path(__file__).parent / "fixtures/related_party_tables.json")
                     .read_text(encoding="utf-8"))["300391_rp_sales"]
    words = [Word(text=t, bbox=BBox(x0=a, y0=b, x1=c, y1=d)) for t, a, b, c, d in fix["words"]]
    tables = extract_note_tables(words, page_index=fix["page"], document_id="t",
                                 source_kind="native", carry_note=(fix["note"], fix["title"]))
    sales = {i.group_hint: str(next(v.value for v in i.values.values() if v.period_label == "current"))
             for t in tables for i in t.items if i.raw_label == "销售商品"}
    assert sales == {"湖北茉莉大药房连锁有限公司": "2046232.67",
                     "湖北好药师舒惠涛社区药房连锁有限公司": "154314.52",
                     "湖北长江大药房连锁武汉有限公司": "-44598.86",
                     "湖北长江大药房连锁有限公司": "-18562.98"}


def test_a_wrapped_caption_reading_only_关联方_opens_no_table():
    """600990: a 坏账准备 movement row captioned "合并范围内 / 关联方" over two lines. Its second line is
    the counterparty word alone; read as a header it split the movement table in two."""
    from app.core.models.geometry import BBox
    from app.services.notes_extract import _opens_a_sub_table
    from app.services.row_reconstruct import Word

    def row(*texts):
        return [Word(text=t, bbox=BBox(x0=0.1 + 0.15 * i, y0=0.5, x1=0.2 + 0.15 * i, y1=0.51))
                for i, t in enumerate(texts)]

    assert not _opens_a_sub_table(row("关联方"))
    assert _opens_a_sub_table(row("关联方", "关联交易内容", "本期发生额", "上期发生额"))


def test_each_item_names_the_counterparty_its_row_is_printed_under(shipped):
    row = NoteItem(raw_label="销售商品", group_hint="湖北茉莉大药房连锁有限公司")
    row.values["k"] = ExtractedValue(basis=Basis.CONSOLIDATED, period_label="current",
                                     value=Decimal("2046232.67"), value_raw=Decimal("2046232.67"),
                                     provenance=Provenance(page_index=244))
    sales = NotesTable(note_number="十四、5", title="出售商品/提供劳务情况表",
                       chapter_title="关联方及关联交易", items=[row], source_pages=[244],
                       source_text="出售商品/提供劳务情况表 湖北茉莉大药房连锁有限公 销售商品 2,046,232.67 司")
    doc = DocumentModel(filename="f.pdf")
    doc.notes = [sales]
    doc.line_items = [_llm_part("sub__rpt_sales", "note 十四、5: 销售商品", "2046232.67", "十四、5", 244)]
    ctx = PipelineContext(settings=get_settings())
    ctx.line_items = shipped
    NoteSourcedStage().run(doc, ctx)
    item = _row(doc, RPT).disclosure["items"][0]
    assert (item["group"], item["caption"], item["note"]) == (
        "湖北茉莉大药房连锁有限公司", "销售商品", "十四、5")
    served = next(r for r in _serialize_rows(doc) if r["canonical_key"] == RPT)
    assert "湖北茉莉大药房连锁有限公司 · 销售商品" in disclosure_text(served)
