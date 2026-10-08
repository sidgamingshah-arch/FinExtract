"""ICON's output is the Company v3 workbook it was built from — filled, not reinterpreted.

The workbook (`app/sample/workbooks/output_csv_icon_v1.xlsx`) is the bank's own layout: a
header block, then one row per line item keyed by the UUID in column A and one value column per
period, with the workbook's own formulas for every total and ratio. `services.export_workbook`
fills a copy of it from a run; `scripts/build_icon_pair.py` writes the map of where each ICON line
sits (`output_csv_icon_v1.json`). These tests hold three things:

* THE MAP IS THE WORKBOOK'S. Every ICON line and the five KPIs have a row, and the row is the one
  the workbook gives that UUID, under the same label.
* WHAT GOES IN, AND WHERE. Figures go into input rows only, in the chosen denomination, oldest
  period first; the header block carries only values the workbook's drop-down lists allow; the
  workbook's formulas are kept, except where the filing prints a total without its parts.
* THE WORKBOOK'S OWN ARITHMETIC AGREES WITH THE SPREAD. Recalculated (`tests/xlsx_eval.py`), the
  filled workbook shows the ICON spread's totals and a balance sheet that balances in both years.
"""
from __future__ import annotations

import datetime as dt
import io
import json
import pathlib

import openpyxl
import pytest

from tests.xlsx_eval import Evaluator

TEMPLATES = pathlib.Path(__file__).resolve().parents[1] / "app" / "sample" / "templates"
WORKBOOKS = TEMPLATES.parent / "workbooks"
P, B, C = "pl_icon__", "bs_icon__", "cov_icon__"
PL, BS, AI = "Profit and Loss Statement", "Balance Sheet", "Additional Information"


@pytest.fixture(scope="module")
def layout():
    return json.loads((WORKBOOKS / "output_csv_icon_v1.json").read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def template():
    return json.loads((TEMPLATES / "output_csv_icon_v1_template.json").read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def line_items():
    from app.schemas.line_items import load_line_item_set

    raw = json.loads((TEMPLATES / "output_csv_icon_line_items.json").read_text(encoding="utf-8"))
    return load_line_item_set(raw, resolve=True)


def _row(key: str, *values) -> dict:
    """A served row: ``values`` are (period, value) or (period, value, period_display)."""
    return {"canonical_key": key, "source_label": key, "values": [
        {"period_label": v[0], "value": str(v[1]), "basis": "standalone",
         "period_display": v[2] if len(v) > 2 else None} for v in values]}


def _fill(template, line_items, rows, **kw) -> openpyxl.Workbook:
    from app.services.export_workbook import build_template_workbook

    kw.setdefault("source_units", {"currency": "", "scale_factor": "100000", "units_label": "lakh"})
    data = build_template_workbook(rows, template, line_item_set=line_items, **kw)
    return openpyxl.load_workbook(io.BytesIO(data))


def _at(book, layout, key: str, period_index: int):
    spot = layout["rows"][key]
    return book[spot["sheet"]].cell(spot["row"], 6 + period_index)


# ── the map ──────────────────────────────────────────────────────────────────────────────────────

def test_every_icon_line_and_kpi_has_the_row_the_workbook_gives_its_uuid(layout, template):
    book = openpyxl.load_workbook(WORKBOOKS / layout["workbook"])
    lines = {}

    def walk(node):
        if isinstance(node, dict):
            if node.get("canonical_key") and node.get("role") != "header":
                lines[node["canonical_key"]] = node["label"]
            for v in node.values():
                walk(v)
        elif isinstance(node, list):
            for v in node:
                walk(v)
    walk(template)
    assert set(lines) == {k for k in layout["rows"] if not k.startswith("kpi:")}
    assert {k for k in layout["rows"] if k.startswith("kpi:")} == {
        f"kpi:{r['key']}" for r in template["kpis"]["ratios"]}
    for key, spot in layout["rows"].items():
        ws = book[spot["sheet"]]
        assert ws.cell(spot["row"], 1).value == spot["uuid"], key
        assert str(ws.cell(spot["row"], 3).value) == spot["sequence"], key
        if key in lines:
            assert " ".join(str(ws.cell(spot["row"], 5).value).split()) == lines[key], key


def test_the_templates_statements_are_the_workbooks_first_three_sheets(layout, template):
    book = openpyxl.load_workbook(WORKBOOKS / layout["workbook"])
    assert layout["sheets"] == [ws.title for ws in book.worksheets[:3]] == [PL, BS, AI]
    assert [s["label"] for s in template["statements"]] == [PL, BS, AI]


# ── what goes in, and where ──────────────────────────────────────────────────────────────────────

def test_the_header_block_says_whose_statements_and_in_what(template, line_items, layout):
    book = _fill(template, line_items, [
        _row(P + "domestic_sales", ("prior", 900), ("current", 1000, "for the year ended 31 March 2025"))],
        entity="SAMPLE INDUSTRIES LIMITED")
    ws = book[PL]
    assert ws["F1"].value == "SAMPLE INDUSTRIES LIMITED"
    assert ws["F6"].value == "Standalone"
    assert (ws["F8"].value, ws["G8"].value) == ("Audited", "Audited")
    assert ws["G9"].value.date() == dt.date(2025, 3, 31)
    # The comparative's own date was not printed, so it is the year before — and says so.
    assert ws["F9"].value.date() == dt.date(2024, 3, 31) and ws["F9"].comment is not None
    assert (ws["F10"].value, ws["G10"].value) == (12, 12)
    assert ws["F11"].value == "INR"                  # the template's own, kept: none was read
    assert ws["F12"].value == "Lakhs"
    # Oldest first: the workbook's growth formula reads the column to the left as the year before.
    assert (_at(book, layout, P + "domestic_sales", 0).value,
            _at(book, layout, P + "domestic_sales", 1).value) == (900, 1000)


def test_a_chosen_denomination_converts_every_figure(template, line_items, layout):
    book = _fill(template, line_items, [_row(P + "domestic_sales", ("current", 1250))],
                 units="crore")
    assert book[PL]["F12"].value == "Crores"
    assert _at(book, layout, P + "domestic_sales", 0).value == pytest.approx(12.5)


def test_an_undeclared_scale_claims_no_denomination(template, line_items, layout):
    book = _fill(template, line_items, [_row(P + "domestic_sales", ("current", 1250))],
                 source_units=None)
    assert book[PL]["F12"].value is None              # not the template's "Crores"
    assert _at(book, layout, P + "domestic_sales", 0).value == 1250


def test_inputs_are_filled_and_the_workbooks_formulas_are_kept(template, line_items, layout):
    book = _fill(template, line_items, [
        _row(P + "domestic_sales", ("current", 1000)), _row(P + "export_sales", ("current", 200)),
        # printed in brackets as a deduction: the workbook subtracts it as printed
        _row(P + "excise_duty", ("current", -50))])
    assert _at(book, layout, P + "excise_duty", 0).value == 50
    assert str(_at(book, layout, P + "gross_sales_total", 0).value).startswith("=ROUND(SUM(")
    assert str(_at(book, layout, P + "net_sales", 0).value).startswith("=ROUND(")
    ev = Evaluator(book)
    assert ev.value(PL, f"F{layout['rows'][P + 'net_sales']['row']}") == 1150


def test_a_total_printed_without_its_parts_keeps_its_printed_figure(template, line_items, layout):
    book = _fill(template, line_items, [
        _row(P + "raw_materials", ("current", 6500)),             # no imported / indigenous split
        _row(B + "receivables", ("current", 2100))])              # no domestic / export split
    raw = _at(book, layout, P + "raw_materials", 0)
    assert raw.value == 6500 and "=ROUND(SUM(F$31:F$32),2)" in raw.comment.text
    # Total Current Assets adds Domestic and Export Receivables, not the Receivables row, so the
    # printed figure goes on the row the line's definition names for an unsplit figure.
    assert str(_at(book, layout, B + "receivables", 0).value).startswith("=")
    dom = _at(book, layout, B + "domestic_receivables", 0)
    assert dom.value == 2100 and dom.comment is not None
    ev = Evaluator(book)
    assert ev.value(BS, f"F{layout['rows'][B + 'total_current_assets']['row']}") == 2100


def test_a_roll_forward_takes_the_printed_balance_and_keeps_its_formula_in_the_comment(
        template, line_items, layout):
    book = _fill(template, line_items, [
        _row(B + "surplus_or_deficit_in_profit_and_loss", ("prior", 800), ("current", 1100)),
        _row(P + "profit_after_tax", ("current", 450))])
    prior, current = (_at(book, layout, B + "surplus_or_deficit_in_profit_and_loss", i)
                      for i in (0, 1))
    assert (prior.value, current.value) == (800, 1100)
    assert "=F65+'Profit and Loss Statement'!G96" in current.comment.text


def test_a_roll_forward_on_a_line_the_filing_does_not_use_is_cleared(template, line_items, layout):
    """Reserves printed as one "Other equity" sit on Other Reserves; left in place, the surplus
    roll-forward would add this year's retained profit to them a second time."""
    book = _fill(template, line_items, [
        _row(B + "other_reserves", ("prior", 3000), ("current", 3450)),
        _row(P + "profit_after_tax", ("prior", 400), ("current", 450))])
    cell = _at(book, layout, B + "surplus_or_deficit_in_profit_and_loss", 1)
    assert cell.value is None and "cleared" in cell.comment.text


def test_the_dividend_rate_is_written_and_the_workbooks_ratios_are_left_to_it(
        template, line_items, layout):
    book = _fill(template, line_items, [
        _row(B + "share_capital", ("current", 100)), _row(P + "equity_dividend", ("current", 20))])
    assert _at(book, layout, "kpi:dividend_rate", 0).value == pytest.approx(20.0)
    assert str(_at(book, layout, "kpi:current_ratio", 0).value).startswith("=IF(")


def test_a_heading_icon_only_derives_is_left_blank(template, line_items, layout):
    """"Investments (other than long term)" is an input row the workbook's totals skip — they add
    its two parts — so a figure there would invite counting them twice."""
    book = _fill(template, line_items, [
        _row(B + "govt_and_trustee_securities", ("current", 100)),
        _row(B + "fixed_deposits_with_banks", ("current", 50))])
    assert _at(book, layout, B + "current_investments", 0).value is None
    assert _at(book, layout, B + "govt_and_trustee_securities", 0).value == 100


def _hidden_columns(ws) -> set[int]:
    return {c for dim in ws.column_dimensions.values() if dim.hidden
            for c in range(dim.min, dim.max + 1)}


def test_only_the_filings_two_years_are_shown(template, line_items):
    """The workbook has thirty period columns; an annual report fills two. The other 28 would show
    the template's zeros and a TRUE balance check, so they are hidden — not removed, so the bank's
    layout and formulas stay as they were."""
    book = _fill(template, line_items, [_row(P + "domestic_sales", ("prior", 900), ("current", 1000))])
    for name in (PL, BS, AI):
        hidden = _hidden_columns(book[name])
        assert hidden == {1, 2, 3} | set(range(8, 36)), name     # A:C is the template's own
        assert book[name].column_dimensions["F"].width == pytest.approx(20.7265625)
    assert str(book[PL]["H23"].value).startswith("=ROUND(SUM(")    # kept, only out of sight


def test_a_filing_with_one_year_shows_one_column(template, line_items):
    book = _fill(template, line_items, [_row(P + "domestic_sales", ("current", 1000))])
    assert _hidden_columns(book[PL]) == {1, 2, 3} | set(range(7, 36))


def test_the_workbook_keeps_its_protection_and_drop_downs(template, line_items):
    book = _fill(template, line_items, [_row(P + "domestic_sales", ("current", 1))])
    assert all(book[name].protection.sheet for name in (PL, BS, AI))
    assert len(book[PL].data_validations.dataValidation) == 7
    assert book.calculation.fullCalcOnLoad


# ── the workbook's own arithmetic, on a filing ───────────────────────────────────────────────────

@pytest.fixture(scope="module")
def schedule_iii_run(line_items):
    pytest.importorskip("fitz")
    pytest.importorskip("reportlab")
    from app.api.routes.extractions import _serialize_rows
    from app.schemas.loader import load_template
    from app.services.documents import run_extraction
    from app.services.working_view import build_working_view
    from tests.test_the_icon_set import _schedule_iii_pdf, _template

    doc, _ctx = run_extraction(_schedule_iii_pdf(title_with_date=True, current_tax="(1)"),
                               filename="schedule_iii.pdf",
                               ontology=build_working_view(line_items),
                               template=load_template(_template()), line_items=line_items)
    return _serialize_rows(doc), doc.unit_context.model_dump(mode="json")


def test_recalculated_the_workbook_shows_the_spread_and_balances(schedule_iii_run, template,
                                                                 line_items, layout):
    from app.services.rollups import figures_as_shown

    rows, units = schedule_iii_run
    book = _fill(template, line_items, rows, source_units=units, entity="SAMPLE INDUSTRIES LIMITED")
    ev = Evaluator(book)
    for i, period in enumerate(("prior", "current")):
        shown = figures_as_shown(template, rows, "standalone", period)
        col = "FG"[i]
        for key in (P + "total_operating_income", P + "profit_before_tax", P + "profit_after_tax",
                    B + "total_current_liabilities", B + "net_worth", B + "total_current_assets",
                    B + "net_block", B + "total_assets", B + "total_liabilities_and_net_worth",
                    B + "tangible_net_worth"):
            got = ev.value(layout["rows"][key]["sheet"], f"{col}{layout['rows'][key]['row']}")
            assert got == pytest.approx(shown[key]), (key, period)
        assert ev.value(BS, f"{col}134") == pytest.approx(0)      # DIFFERENCE IN B/S
        assert ev.value(BS, f"{col}133") is True


# ── the endpoint ─────────────────────────────────────────────────────────────────────────────────

@pytest.fixture()
def icon_run(client):
    from app.db.base import SessionLocal
    from app.db.models import Document, ExtractionRun, TemplateVersion

    with SessionLocal() as session:
        tv = (session.query(TemplateVersion).filter_by(template_key="output_csv_icon_v1")
              .order_by(TemplateVersion.version.desc()).first())
        assert tv is not None, "the shipped ICON template is seeded"
        doc = Document(filename="icon-workbook.pdf", content_hash="icon-workbook-1",
                       object_key="tests/icon-workbook.pdf")
        session.add(doc)
        session.flush()
        run = ExtractionRun(
            document_id=doc.id, run_number=1, status="succeeded",
            options={"template_version_id": tv.id},
            result={"entity": "SAMPLE INDUSTRIES LIMITED",
                    "units": {"currency": "INR", "scale_factor": "100000", "units_label": "lakh"},
                    "rows": [_row(P + "domestic_sales", ("current", 1000)),
                             _row(B + "share_capital", ("current", 100))]})
        plain = Document(filename="no-workbook.pdf", content_hash="no-workbook-1",
                         object_key="tests/no-workbook.pdf")
        session.add(plain)
        session.add(run)
        session.flush()
        session.add(ExtractionRun(document_id=plain.id, run_number=1, status="succeeded",
                                  result={"rows": [_row(P + "domestic_sales", ("current", 1))]}))
        session.commit()
        return doc.id, plain.id


def test_the_endpoint_serves_the_filled_workbook(client, icon_run, layout):
    doc_id, plain_id = icon_run
    assert client.get(f"/api/v1/documents/{doc_id}/run").json()["workbook_export"] is True
    got = client.get(f"/api/v1/documents/{doc_id}/export?fmt=workbook")
    assert got.status_code == 200, got.text
    assert 'filename="icon-workbook - ICON.xlsx"' in got.headers["content-disposition"]
    book = openpyxl.load_workbook(io.BytesIO(got.content))
    assert book[PL]["F1"].value == "SAMPLE INDUSTRIES LIMITED"
    assert _at(book, layout, P + "domestic_sales", 0).value == 1000

    assert client.get(f"/api/v1/documents/{plain_id}/run").json()["workbook_export"] is False
    assert client.get(f"/api/v1/documents/{plain_id}/export?fmt=workbook").status_code == 404
