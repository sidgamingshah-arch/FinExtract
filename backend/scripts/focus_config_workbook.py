#!/usr/bin/env python
"""THE 8 FOCUS FAMILIES AS AN EDITABLE WORKBOOK — values, not just controls, with a README.

WHY A SECOND WORKBOOK. `line_item_options_workbook.py` answers "what does each CONTROL do and does
anyone use it" — a question about the screen. This answers "what does each LINE ITEM currently say,
and is that value any good" — a question about the configuration, for an analyst who wants to work
offline on the vocabulary rather than click through 77 forms.

WHAT MAKES IT USEFUL RATHER THAN A DUMP:

  * ONE SHEET PER FAMILY, one row per part, one column per field. A family is the unit an analyst
    reasons about, because a cascade's rungs compete for the same notes.
  * NOTE-HEADER HINTS AND ROW HINTS ARE IN SEPARATE COLUMN BLOCKS, shaded apart, because mixing
    them is the measured failure — a heading names the container and a row names the content, and
    139 of the 144 flagged values in this set are note-level values written as row content.
  * EVERY VALUE CARRIES ITS MEASURED REACH: how many of the corpus filings its patterns actually
    claim a heading in. A value reaching 1 of 18 is a constant wearing a config field, and that is
    the number that tells an author where to spend their time.
  * `definition` IS EXPORTED READ-ONLY and greyed. It is the analyst's to write and nothing here
    should imply otherwise.

ROUND TRIP. The workbook is the same shape `--import` reads back, so an edited file becomes a JSON
patch for `POST /api/v1/line-items` ("the only door into the store") without hand-transcription.
A list cell is one value per line inside the cell — never comma-separated, because a regex may
contain a comma and a caption certainly can.

    python scripts/focus_config_workbook.py -o ../_vocab/focus_config.xlsx
    python scripts/focus_config_workbook.py --import ../_vocab/focus_config.xlsx -o patch.json
"""
from __future__ import annotations

import argparse
import json
import pathlib
import re
import sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

SEED = (pathlib.Path(__file__).resolve().parent.parent / "app" / "sample" / "templates"
        / "output_csv_hk_line_items.json")
VOCAB = pathlib.Path(__file__).resolve().parent.parent.parent / "_vocab"

FAMILIES = [
    ("Revenue", "is_pl__sales_revenues"),
    ("Deprec Oper Exp", "is_pl__deprec_and_impairment_oper_exp"),
    ("Deprec COS", "is_pl__deprec_and_impairment_cos"),
    ("Fincl Assets CP", "bs_ca__secur_and_other_fincl_assets_cp"),
    ("Fincl Assets LTP", "bs_nca__secur_and_other_fincl_assets_ltp"),
    ("Due from Rel Parties", "bs_nca__due_from_related_parties_ltp"),
    ("Other Receivables CP", "bs_ca__other_receivables_cp"),
    ("Contingent Liabs", "notes__contingent_liabilities"),
]

# The two hint levels, kept apart on purpose — see the module docstring.
NOTE_HINTS = ["note_title_any", "note_terms"]
ROW_HINTS = ["row_caption_any", "row_caption_none", "row_terms", "row_terms_none"]
PROSE_HINTS = ["prose_any"]
CRITERIA = ["include_criteria", "exclude_criteria"]
READ_ONLY = ["key", "label", "parent", "definition", "reach"]

FIELD_HELP = {
    "key": "The line item's identity. NEVER EDIT — a cascade rung names its parts by this, and "
           "renaming one silently empties the rung.",
    "label": "What the line is called on screen and in the export.",
    "parent": "The derived line this part feeds. Read-only here; it is the cascade's wiring.",
    "definition": "READ-ONLY IN THIS WORKBOOK, and yours to write. Prose describing what the line "
                  "MEANS. Read by the description-matching tier and, for an extracted line, by the "
                  "model. Not a pattern — never write regex here.",
    "reach": "MEASURED, not authored: how many of the corpus filings this part's `note_title_any` "
             "claims a note heading in. 0 means nothing can reach it; 1 means the pattern fits one "
             "filing and is a constant wearing a config field.",
    "note_title_any": "NOTE-HEADER HINT. REGEXES over a note's HEADING — which note this line's "
                      "figure lives in. A heading names the CONTAINER: 'ADMINISTRATIVE EXPENSES', "
                      "或有負債, 重大担保. Matched with re.search, case-insensitive, against the "
                      "heading AS PRINTED. Beware false friends: 'GUARANTEED NOTES' is a debt "
                      "instrument, not a contingency.",
    "note_terms": "NOTE-HEADER HINT. Plain WORDS, not regexes, scored against note HEADINGS by "
                  "IDF-weighted cosine. Same job as the row above by a different mechanism: a "
                  "phrasing nobody wrote a pattern for still ranks instead of not firing. KEEP "
                  "THESE ABOUT THE NOTE, not the row — a note-level field full of row-content "
                  "words scores 0.000 against the very heading its line belongs to.",
    "row_caption_any": "ROW HINT. REGEXES over a ROW CAPTION inside the chosen note — which rows "
                       "count toward this figure. A row names the CONTENT: 'Depreciation of fixed "
                       "assets', 固定资产折旧. A row the note prints but no pattern here matches "
                       "contributes nothing.",
    "row_caption_none": "ROW HINT, VETO. Rows that must NOT count even if a pattern above matched. "
                        "Each is compiled on save — a torn veto stops vetoing silently and the "
                        "wrong rows get summed.",
    "row_terms": "ROW HINT. Plain words scored against ROW CAPTIONS. These also reach the MODEL "
                 "and gate its answers, so a term here does double duty: it ranks a row and tells "
                 "the model what the line is called.",
    "row_terms_none": "ROW HINT, VETO, scored. A row whose caption carries one of these is not "
                      "this line's however well it otherwise scores.",
    "prose_any": "SENTENCE HINT — a third level, neither heading nor row. REGEXES over a note's "
                 "SENTENCES, for a figure the filing states in words and tabulates nowhere. "
                 "Authored for SENTENCE length: the reference footnote has 87 characters between "
                 "its two subjects, so caption-length bounds like .{0,40} do not reach. EMPTY "
                 "MEANS NO PROSE ROUTE, which is the safe default.",
    "include_criteria": "Prose criteria the description tier and the model read — what COUNTS as "
                        "this line. Not patterns.",
    "exclude_criteria": "Prose criteria — what does NOT count, however close the wording.",
}

README = [
    ("HOW TO USE THIS WORKBOOK", ""),
    ("", "One sheet per focus family, one row per part. Edit the value cells; leave the grey "
         "read-only columns alone."),
    ("", "A LIST CELL HOLDS ONE VALUE PER LINE inside the cell (alt+enter). Never comma-separate: "
         "a regex may contain a comma and a caption certainly can."),
    ("", "An EMPTY cell means an empty list — a configured empty, which is a real declaration and "
         "not 'use a default'."),
    ("", "Hand the edited file back and it becomes a patch for POST /api/v1/line-items, which is "
         "the only door into the store and runs the same checks the shipped file passes at boot."),
    ("", ""),
    ("THE ONE RULE THAT MATTERS MOST", ""),
    ("", "NOTE-HEADER HINTS AND ROW HINTS ARE DIFFERENT VOCABULARIES AND MUST NOT BE MIXED."),
    ("", "A note's heading names the CONTAINER — 'ADMINISTRATIVE EXPENSES', 管理费用, 或有負債."),
    ("", "A row's caption names the CONTENT — 'Depreciation of fixed assets', 固定资产折旧."),
    ("", "Measured on this set: a blended probe scored 0.000 against the very heading its line "
         "belongs to, because every one of its tokens was a content word. A value written at the "
         "wrong level does not degrade gracefully — it degrades to silence."),
    ("", "Measured on the 77 parts: 144 values are mixed, ambiguous or dead, and 139 of those 144 "
         "are note-level fields written as row content."),
    ("", ""),
    ("WHAT `reach` MEANS", ""),
    ("", "How many corpus filings this part's note_title_any claims a heading in. It is measured, "
         "not authored, and it is the number to spend time on:"),
    ("", "  0  nothing can reach this part at all"),
    ("", "  1  the pattern fits one filing — a hard-coded constant wearing a config field"),
    ("", "  high  the vocabulary generalises"),
    ("", ""),
    ("FIELD BY FIELD", ""),
]


def _reach(item: dict, corpus: dict[str, list[str]]) -> int:
    pats = list(((item.get("note_source") or {}).get("note_title_any")) or [])
    rxs = []
    for p in pats:
        try:
            rxs.append(re.compile(p, re.I))
        except re.error:
            continue
    if not rxs:
        return 0
    return sum(1 for heads in corpus.values() if any(rx.search(h) for h in heads for rx in rxs))


def _corpus() -> dict[str, list[str]]:
    out: dict[str, list[str]] = {}
    for path in sorted(VOCAB.glob("*.json")):
        if path.name == "INDEX.json":
            continue
        out[path.stem] = sorted({(n.get("title") or "").strip()
                                 for n in json.loads(path.read_text(encoding="utf-8"))
                                 .get("notes") or [] if (n.get("title") or "").strip()})
    return out


def export(out_path: pathlib.Path) -> int:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    raw = json.loads(SEED.read_text(encoding="utf-8"))
    items = [i for i in raw["items"]]
    corpus = _corpus()

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ink = "1F2937"
    head_font = Font(bold=True, color="FFFFFF", size=10)
    note_fill = PatternFill("solid", fgColor="1E3A5F")      # note-header hints
    row_fill = PatternFill("solid", fgColor="4A1E5F")       # row hints
    prose_fill = PatternFill("solid", fgColor="5F3A1E")     # sentence hints
    ro_fill = PatternFill("solid", fgColor="6B7280")        # read-only
    crit_fill = PatternFill("solid", fgColor="1E5F4A")      # criteria
    grey = PatternFill("solid", fgColor="F3F4F6")
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(bottom=Side(style="thin", color="D1D5DB"))

    # ── README ────────────────────────────────────────────────────────────────────────────────
    ws = wb.create_sheet("README")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 118
    r = 1
    for a, b in README:
        ca, cb = ws.cell(r, 1, a), ws.cell(r, 2, b)
        if a:
            ca.font = Font(bold=True, color=ink, size=11)
        cb.alignment = wrap
        cb.font = Font(size=10, color=ink)
        r += 1
    for field in READ_ONLY + NOTE_HINTS + ROW_HINTS + PROSE_HINTS + CRITERIA:
        ws.cell(r, 1, field).font = Font(bold=True, color=ink, size=10, name="Consolas")
        c = ws.cell(r, 2, FIELD_HELP.get(field, ""))
        c.alignment = wrap
        c.font = Font(size=10, color=ink)
        ws.row_dimensions[r].height = 30
        r += 1
    ws.freeze_panes = "A2"

    # ── one sheet per family ──────────────────────────────────────────────────────────────────
    cols = READ_ONLY + NOTE_HINTS + ROW_HINTS + PROSE_HINTS + CRITERIA
    fill_of = {**{c: ro_fill for c in READ_ONLY},
               **{c: note_fill for c in NOTE_HINTS},
               **{c: row_fill for c in ROW_HINTS},
               **{c: prose_fill for c in PROSE_HINTS},
               **{c: crit_fill for c in CRITERIA}}

    for sheet_name, parent in FAMILIES:
        kids = [i for i in items if i.get("parent") == parent]
        ws = wb.create_sheet(sheet_name[:31])
        ws.cell(1, 1, f"{parent}   —   {len(kids)} parts").font = Font(bold=True, size=11,
                                                                       color=ink)
        for j, col in enumerate(cols, start=1):
            c = ws.cell(2, j, col)
            c.font = head_font
            c.fill = fill_of[col]
            c.alignment = Alignment(wrap_text=True, vertical="bottom")
            ws.column_dimensions[c.column_letter].width = (
                14 if col in ("key", "reach") else 34 if col in NOTE_HINTS + ROW_HINTS else 26)
        ws.column_dimensions["A"].width = 46
        ws.column_dimensions["D"].width = 50

        for i, kid in enumerate(kids, start=3):
            ns = kid.get("note_source") or {}
            for j, col in enumerate(cols, start=1):
                if col == "reach":
                    v: object = _reach(kid, corpus)
                elif col in NOTE_HINTS + ROW_HINTS + PROSE_HINTS:
                    v = "\n".join(str(x) for x in (ns.get(col) or []))
                elif col in ("include_criteria", "exclude_criteria"):
                    v = "\n".join(str(x) for x in (kid.get(col) or []))
                else:
                    v = kid.get(col) or ""
                c = ws.cell(i, j, v)
                c.alignment = wrap
                mono = col in (NOTE_HINTS + ROW_HINTS + PROSE_HINTS + ["key"])
                c.font = Font(size=9, name="Consolas" if mono else "Calibri")
                c.border = thin
                if col in READ_ONLY:
                    c.fill = grey
            ws.row_dimensions[i].height = 58
        ws.freeze_panes = "B3"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"wrote {out_path}")
    print(f"  {len(FAMILIES)} family sheets + README, "
          f"{sum(1 for i in items if i.get('parent') in dict((p, n) for n, p in FAMILIES))} parts")
    print(f"  reach measured against {len(corpus)} corpus filings")
    return 0


def import_(xlsx: pathlib.Path, out_json: pathlib.Path) -> int:
    """An edited workbook back to a JSON patch — one entry per CHANGED value, nothing else.

    ONLY CHANGES, because a full re-export would rewrite 77 items and make a two-value edit
    indistinguishable from a rewrite in review.
    """
    import openpyxl

    raw = json.loads(SEED.read_text(encoding="utf-8"))
    items = {i["key"]: i for i in raw["items"]}
    wb = openpyxl.load_workbook(xlsx)
    changes: list[dict] = []

    for sheet_name, _parent in FAMILIES:
        if sheet_name[:31] not in wb.sheetnames:
            continue
        ws = wb[sheet_name[:31]]
        header = [c.value for c in ws[2]]
        for row in ws.iter_rows(min_row=3):
            cells = {header[j]: row[j].value for j in range(len(header)) if header[j]}
            key = str(cells.get("key") or "").strip()
            if not key or key not in items:
                continue
            ns = items[key].get("note_source") or {}
            for col in NOTE_HINTS + ROW_HINTS + PROSE_HINTS + CRITERIA:
                if col not in cells:
                    continue
                new = [l.strip() for l in str(cells[col] or "").splitlines() if l.strip()]
                old = [str(x) for x in ((ns.get(col) if col not in CRITERIA
                                         else items[key].get(col)) or [])]
                if new != old:
                    changes.append({"key": key,
                                    "field": (col if col in CRITERIA
                                              else f"note_source.{col}"),
                                    "was": old, "now": new})
    out_json.write_text(json.dumps({"changes": changes}, ensure_ascii=False, indent=1),
                        encoding="utf-8")
    print(f"{len(changes)} changed value(s) -> {out_json}")
    for ch in changes[:12]:
        print(f"   {ch['key']:48s} {ch['field']:28s} {len(ch['was'])} -> {len(ch['now'])}")
    if changes:
        print("\nCheck them before publishing:")
        print(f"  python scripts/check_proposals.py --proposals {out_json}")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("-o", "--out", type=pathlib.Path,
                    default=VOCAB / "focus_config.xlsx")
    ap.add_argument("--import", dest="imp", type=pathlib.Path, default=None)
    args = ap.parse_args()
    if args.imp:
        return import_(args.imp, args.out if args.out.suffix == ".json"
                       else args.imp.with_suffix(".patch.json"))
    return export(args.out)


if __name__ == "__main__":
    raise SystemExit(main())
