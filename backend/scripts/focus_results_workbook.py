#!/usr/bin/env python
"""The eight focus concepts across both filings, with where every figure came from.

Reads the JSON that `run_focus_filings.py` writes for each filing and lays the results out for
review: the figure, the cascade rung or notes it came from, every contribution with its note and
caption, and — for a concept that came out empty — WHICH KIND of empty it is.

THAT LAST COLUMN IS THE POINT. A concept empty because the filing does not disclose it, one empty
because no cascade rung resolved, and one empty because its configuration cannot reach it at all are
three different facts. They look identical in a spreadsheet of figures, and only two of them are
something to act on.

    python scripts/focus_results_workbook.py run_a.json run_b.json -o results.xlsx
"""
from __future__ import annotations

import argparse
import json
import pathlib


def build(runs: list[dict], out: pathlib.Path) -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    HEAD = PatternFill("solid", fgColor="1F3864")
    HEAD_F = Font(color="FFFFFF", bold=True, size=10)
    GOT = PatternFill("solid", fgColor="E2EFDA")
    EMPTY = PatternFill("solid", fgColor="FCE4E4")
    BLOCKED = PatternFill("solid", fgColor="FFF2CC")
    THIN = Side(style="thin", color="D0D0D0")
    BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    TOP = Alignment(vertical="top", wrap_text=True)
    NUM = Alignment(vertical="top", horizontal="right")
    MONO = Font(name="Consolas", size=9)

    wb = Workbook()

    def sheet(title, cols, first=False):
        ws = wb.active if first else wb.create_sheet(title)
        ws.title = title
        ws.append([c[0] for c in cols])
        for n, (_h, w) in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(n)].width = w
            c = ws.cell(row=1, column=n)
            c.fill, c.font, c.border = HEAD, HEAD_F, BOX
            c.alignment = Alignment(vertical="center", wrap_text=True, horizontal="center")
        ws.row_dimensions[1].height = 34
        ws.freeze_panes = "B2"
        return ws

    def put(ws, values, *, fill=None, mono=(), nums=()):
        ws.append(values)
        r = ws.max_row
        for c in range(1, len(values) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border, cell.alignment = BOX, (NUM if c in nums else TOP)
            if fill:
                cell.fill = fill
            if c in mono:
                cell.font = MONO
        return r

    keys = [c["key"] for c in runs[0]["concepts"]]
    labels = {c["key"]: c["label"] for c in runs[0]["concepts"]}

    # ── 1. the answer ─────────────────────────────────────────────────────────────────────────
    cols = [("Concept", 42), ("Label", 30)]
    for r in runs:
        name = r["filing"].replace(".pdf", "")
        cols += [(f"{name}\ncurrent", 17), (f"{name}\nprior", 17), (f"{name}\nsource", 26),
                 (f"{name}\nif empty, why", 40)]
    ws = sheet("1. Figures", cols, first=True)

    for key in keys:
        vals = [key, labels.get(key, "")]
        fills = []
        for r in runs:
            c = next(x for x in r["concepts"] if x["key"] == key)
            per = {f["period"]: (f["value"] or f["text"]) for f in c["figures"]}
            src = ""
            if c["trail"]:
                methods = sorted({t["method"] for t in c["trail"]})
                src = ", ".join(methods)
            elif per:
                src = "face / matcher"
            vals += [per.get("current", ""), per.get("prior", ""), src, c["why_empty"]]
            fills.append(GOT if per else (BLOCKED if "refuses" in c["why_empty"] else EMPTY))
        row = put(ws, vals, nums=tuple(range(3, len(vals) + 1, 4)))
        for i, f in enumerate(fills):
            for c in range(3 + i * 4, 7 + i * 4):
                ws.cell(row=row, column=c).fill = f

    # ── 2. every contribution ─────────────────────────────────────────────────────────────────
    ts = sheet("2. Trails", [
        ("Filing", 16), ("Concept", 40), ("Rung / method", 24), ("Formula", 44),
        ("Contribution", 44), ("Value", 17), ("Counted", 9), ("Deducted", 9), ("Note", 9)])
    for r in runs:
        for c in r["concepts"]:
            for t in c["trail"]:
                if not t["inputs"]:
                    put(ts, [r["filing"], c["key"], t["method"], t["formula"] or "",
                             "(no inputs recorded)", t["result"], "", "", ""])
                for i in t["inputs"]:
                    put(ts, [r["filing"], c["key"], t["method"], t["formula"] or "",
                             i["label"] or "", i["value"],
                             "yes" if i["counted"] else "no",
                             "yes" if i["deducted"] else "", i.get("note") or ""],
                        mono=(5,), nums=(6,))

    # ── 2b. deterministic vs LLM, per concept ────────────────────────────────────────────────
    #
    # WHY THIS SHEET EXISTS. It is what found the defect worth knowing about: on laisun.pdf the
    # deterministic tiers map the face revenue to 4,995,768 and the LLM run does not map it at all,
    # so a low-precedence segment rung filled the line with 2,609,259 instead. The model is being
    # asked about rows the alias tier already answered correctly, and answering worse. A single
    # column of "the figure" would have hidden that.
    pairs: dict[str, dict[str, dict]] = {}
    for r in runs:
        name = r["filing"].replace(".pdf", "")
        mode = "llm" if (r.get("llm_calls") or 0) else "deterministic"
        pairs.setdefault(name, {})[mode] = r
    both = {n: m for n, m in pairs.items() if len(m) == 2}
    if both:
        cs = sheet("2b. Deterministic vs LLM", [
            ("Filing", 16), ("Concept", 42), ("Deterministic", 18), ("With LLM", 18),
            ("Agree?", 12), ("What it means", 56)])
        for name, modes in sorted(both.items()):
            det, llm = modes["deterministic"], modes["llm"]
            by = {c["key"]: c for c in llm["concepts"]}
            for c in det["concepts"]:
                o = by.get(c["key"], {})
                a = next((f["value"] for f in c["figures"] if f["period"] == "current"), None)
                b = next((f["value"] for f in (o.get("figures") or [])
                          if f["period"] == "current"), None)
                if a is None and b is None:
                    verdict, meaning, fill = "both empty", "", None
                elif a == b:
                    verdict, meaning, fill = "same", "", GOT
                elif a is not None and b is None:
                    verdict, meaning, fill = ("LLM LOST it",
                        "the deterministic tiers reached this concept and the LLM run did not — "
                        "the model was asked about a row the alias tier already answered", EMPTY)
                elif a is None:
                    verdict, meaning, fill = ("LLM found it",
                        "only the model reached this concept", GOT)
                else:
                    verdict, meaning, fill = ("DISAGREE",
                        "two different figures for one line. The deterministic one comes off the "
                        "face; the LLM one came from a cascade rung after the face mapping was "
                        "lost. Trust the face.", EMPTY)
                put(cs, [name, c["key"], a or "", b or "", verdict, meaning],
                    fill=fill, nums=(3, 4))

    # ── 3. the runs themselves ────────────────────────────────────────────────────────────────
    rs = sheet("3. Runs", [("Filing", 20), ("Pages", 8), ("Notes", 8), ("Rows", 8),
                            ("Seconds", 9), ("LLM calls", 10), ("Tokens in", 12),
                            ("Tokens out", 12), ("Concepts filled", 14)])
    for r in runs:
        filled = sum(1 for c in r["concepts"] if c["figures"])
        put(rs, [r["filing"], r["pages"], r["notes"], r["line_items"], r["seconds"],
                 r.get("llm_calls"), r.get("llm_tokens_in"), r.get("llm_tokens_out"),
                 f"{filled} of {len(r['concepts'])}"],
            nums=(2, 3, 4, 5, 6, 7, 8))

    # ── 4. the pipeline's own account ─────────────────────────────────────────────────────────
    ls = sheet("4. Log", [("Filing", 18), ("Line", 118)])
    for r in runs:
        for line in r.get("log") or []:
            put(ls, [r["filing"], line], mono=(2,))

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"wrote {out}")
    for r in runs:
        filled = sum(1 for c in r["concepts"] if c["figures"])
        print(f"  {r['filing']:18s} {filled} of {len(r['concepts'])} focus concepts filled "
              f"({r['pages']} pages, {r['seconds']}s)")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("runs", nargs="+")
    ap.add_argument("-o", "--out", default="focus_results.xlsx")
    args = ap.parse_args()
    runs = [json.loads(pathlib.Path(p).read_text(encoding="utf-8")) for p in args.runs]
    return build(runs, pathlib.Path(args.out).resolve())


if __name__ == "__main__":
    raise SystemExit(main())
