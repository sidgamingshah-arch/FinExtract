from __future__ import annotations

import json
import re
import sqlite3
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl

from app.schemas.loader import load_ontology
from app.services.ontology_xlsx import build_ontology_xlsx


@dataclass
class LegacyRow:
    field_id: str
    section: str
    target_field: str
    canonical_full_form: str
    mapping_method: str
    row_no: int


def _norm(text: str) -> str:
    t = (text or "").strip().lower()
    t = re.sub(r"[^a-z0-9]+", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def _read_latest_ontology(backend_dir: Path) -> dict[str, Any]:
    db_path = backend_dir / "finex.db"
    if db_path.exists():
        con = sqlite3.connect(db_path)
        cur = con.cursor()
        cur.execute(
            """
            select definition
            from ontology_versions
            where ontology_key = 'output_csv_hk'
            order by datetime(created_at) desc
            limit 1
            """
        )
        row = cur.fetchone()
        if row and row[0]:
            raw = row[0]
            return json.loads(raw) if isinstance(raw, str) else raw
    sample = backend_dir / "app" / "sample" / "templates" / "output_csv_hk_ontology.json"
    return json.loads(sample.read_text(encoding="utf-8"))


def _read_legacy_rows(xlsx_path: Path) -> tuple[list[LegacyRow], dict[str, dict[str, str]]]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["Field Ontology"]
    rows: list[LegacyRow] = []

    for r in range(4, ws.max_row + 1):
        field_id = str(ws.cell(r, 1).value or "").strip()
        if not field_id:
            continue
        section = str(ws.cell(r, 4).value or "").strip()
        target_field = str(ws.cell(r, 5).value or "").strip()
        canonical_full = str(ws.cell(r, 6).value or "").strip()
        mapping_method = str(ws.cell(r, 8).value or "").strip().upper()
        rows.append(
            LegacyRow(
                field_id=field_id,
                section=section,
                target_field=target_field,
                canonical_full_form=canonical_full,
                mapping_method=mapping_method,
                row_no=r,
            )
        )

    ws2 = wb["Computed Field Logic"]
    compute_meta: dict[str, dict[str, str]] = {}
    for r in range(4, ws2.max_row + 1):
        field_id = str(ws2.cell(r, 1).value or "").strip()
        if not field_id:
            continue
        compute_meta[field_id] = {
            "compute_type": str(ws2.cell(r, 5).value or "").strip().upper(),
            "deps": str(ws2.cell(r, 7).value or "").strip(),
            "rule": str(ws2.cell(r, 8).value or "").strip(),
        }

    return rows, compute_meta


def _index_current_mappings(defn: dict[str, Any]) -> tuple[dict[str, dict[str, Any]], dict[str, str]]:
    mappings = defn.get("mappings") or []
    by_key = {m.get("canonical_key"): m for m in mappings if m.get("canonical_key")}
    text_to_key: dict[str, str] = {}

    for m in mappings:
        key = m.get("canonical_key")
        if not key:
            continue
        candidates = [m.get("label", ""), key]
        candidates.extend(m.get("aliases") or [])
        for txt in candidates:
            n = _norm(str(txt))
            if n and n not in text_to_key:
                text_to_key[n] = key

    return by_key, text_to_key


def _best_match(row: LegacyRow, text_to_key: dict[str, str]) -> str | None:
    for cand in (row.target_field, row.canonical_full_form):
        n = _norm(cand)
        if n in text_to_key:
            return text_to_key[n]
    return None


def _append_unique(lst: list[str], value: str) -> None:
    v = value.strip()
    if not v:
        return
    if v not in lst:
        lst.append(v)


def _parse_dep_ids(deps_text: str) -> list[str]:
    return re.findall(r"\b[A-Z]{2,10}_\d{3}\b", deps_text or "")


def merge_legacy_into_current(backend_dir: Path, legacy_xlsx: Path) -> tuple[dict[str, Any], dict[str, Any]]:
    current = _read_latest_ontology(backend_dir)
    rows, compute_meta = _read_legacy_rows(legacy_xlsx)
    by_key, text_to_key = _index_current_mappings(current)

    field_id_to_key: dict[str, str] = {}
    unmatched: list[LegacyRow] = []
    matched_rows: list[tuple[LegacyRow, str]] = []

    for row in rows:
        key = _best_match(row, text_to_key)
        if not key:
            unmatched.append(row)
            continue
        field_id_to_key[row.field_id] = key
        matched_rows.append((row, key))

    first_order: dict[str, int] = {}
    by_method: dict[str, int] = defaultdict(int)

    for row, key in matched_rows:
        by_method[row.mapping_method] += 1
        if key not in first_order or row.row_no < first_order[key]:
            first_order[key] = row.row_no

        m = by_key[key]
        m.setdefault("aliases", [])
        _append_unique(m["aliases"], row.target_field)
        _append_unique(m["aliases"], row.canonical_full_form)

        if row.mapping_method in {"AUTO_EXTRACT", "AUTO_CLASSIFY"}:
            m["extraction_mode"] = "extract"

        if row.mapping_method in {"COMPUTE_TOTAL", "COMPUTE_RESIDUAL"}:
            m["alias_matching"] = "disabled"
            m["extraction_mode"] = "extract_or_derive"
            meta = compute_meta.get(row.field_id, {})
            rule = meta.get("rule", "")
            if rule:
                m["derivation"] = rule

            deps = _parse_dep_ids(meta.get("deps", ""))
            dep_keys = [field_id_to_key[d] for d in deps if d in field_id_to_key and field_id_to_key[d] != key]
            dep_keys = list(dict.fromkeys(dep_keys))

            if row.mapping_method == "COMPUTE_TOTAL":
                if dep_keys:
                    m["children_if_decomposed"] = dep_keys

            if row.mapping_method == "COMPUTE_RESIDUAL":
                m["value_scope"] = "exclusive_residual"
                m["extraction_mode"] = "derive"
                if dep_keys:
                    m["expected_components"] = dep_keys
                section = (m.get("section_scope") or [""])[0] if isinstance(m.get("section_scope"), list) else ""
                m["residual_policy"] = {
                    "framework": "residual_framework",
                    "section_scope": section,
                    "population": "sweep_only",
                    "cross_section": False,
                    "notes_as_source": False,
                    "plug": False,
                    "itemise": True,
                }

    # Preserve extraction order by assigning descending priorities from first appearance.
    ordered = sorted(first_order.items(), key=lambda kv: kv[1])
    max_p, min_p = 95, 25
    span = max(1, len(ordered) - 1)
    for i, (key, _row_no) in enumerate(ordered):
        priority = int(round(max_p - (max_p - min_p) * (i / span)))
        by_key[key]["match_priority"] = priority

    current.setdefault("metadata", {})
    current["metadata"].setdefault("changes", [])
    current["metadata"]["changes"].append(
        "Merged HK v1.6 workbook logic: aliases, compute totals/residuals, and extraction order priorities."
    )

    # Validate merged ontology structure.
    load_ontology(current)

    report = {
        "legacy_rows": len(rows),
        "matched_rows": len(matched_rows),
        "unmatched_rows": len(unmatched),
        "method_counts_matched": dict(by_method),
        "unmatched_sample": [
            {
                "field_id": r.field_id,
                "target_field": r.target_field,
                "canonical_full_form": r.canonical_full_form,
                "mapping_method": r.mapping_method,
            }
            for r in unmatched[:40]
        ],
        "unmatched_all": [
            {
                "field_id": r.field_id,
                "section": r.section,
                "target_field": r.target_field,
                "canonical_full_form": r.canonical_full_form,
                "mapping_method": r.mapping_method,
                "row_no": r.row_no,
            }
            for r in unmatched
        ],
    }
    return current, report


def main() -> None:
    backend_dir = Path(__file__).resolve().parents[1]
    legacy_xlsx = backend_dir / "_exports" / "China_HongKong_Financial_Extraction_Ontology_v1.6.xlsx"
    out_json = backend_dir / "_exports" / "output_csv_hk_v16_merged.json"
    out_xlsx = backend_dir / "_exports" / "output_csv_hk_v16_merged.xlsx"
    out_report = backend_dir / "_exports" / "output_csv_hk_v16_merge_report.json"

    merged, report = merge_legacy_into_current(backend_dir, legacy_xlsx)

    out_json.write_text(json.dumps(merged, indent=2, ensure_ascii=False), encoding="utf-8")
    out_xlsx.write_bytes(build_ontology_xlsx(merged, filename_hint="output_csv_hk_v16_merged"))
    out_report.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")

    print("MERGE_OK=1")
    print("OUT_JSON=", out_json)
    print("OUT_XLSX=", out_xlsx)
    print("OUT_REPORT=", out_report)
    print("MATCHED=", report["matched_rows"], "/", report["legacy_rows"])
    print("UNMATCHED=", report["unmatched_rows"])


if __name__ == "__main__":
    main()
